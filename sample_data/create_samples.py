import os
import sys

# Ensure root directory is on Python path
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import docx
from services.pdf_export import generate_pdf_resume

SAMPLE_RESUME_DATA = {
    "contact": {
        "name": "Alex Kumar",
        "email": "alex.kumar@example.com",
        "phone": "+1 (555) 234-5678",
        "github": "github.com/alexkumar-dev",
        "linkedin": "linkedin.com/in/alexkumar-dev"
    },
    "summary": "Motivated Computer Science undergraduate with hands-on experience in full-stack web development and applied machine learning. Proven record building accessible web applications and data-driven prediction models with verifiable evidence.",
    "skills": {
        "Programming Languages": ["Python", "JavaScript", "SQL", "HTML5", "CSS3"],
        "Frontend & Frameworks": ["HTML", "CSS", "JavaScript", "React", "Bootstrap"],
        "Backend & Tools": ["Node.js", "Git", "GitHub", "SQLite"],
        "Machine Learning": ["Machine Learning", "Scikit-Learn", "Pandas", "NumPy"]
    },
    "projects": [
        {
            "title": "Movie Success Prediction System",
            "technologies": "Python, Machine Learning, Scikit-Learn, Pandas",
            "description": "Engineered an end-to-end predictive pipeline analyzing historical box office metrics and sentiment scores to forecast opening weekend film revenue.",
            "github_url": "https://github.com/alexkumar-dev/movie-success-pred"
        },
        {
            "title": "Student Portfolio & Campus Portal",
            "technologies": "HTML, CSS, JavaScript, React, Git",
            "description": "Designed and deployed a responsive web portal allowing campus students to track study clubs, browse events, and share peer notes.",
            "github_url": "https://github.com/alexkumar-dev/campus-portal"
        },
        {
            "title": "Air Quality Prediction Model",
            "technologies": "Python, Machine Learning, NumPy, Scikit-Learn",
            "description": "Built regression model to predict particulate matter air quality index (PM2.5) across urban sensor stations using Python.",
            "github_url": "https://github.com/alexkumar-dev/air-quality-ml"
        }
    ],
    "education": [
        {
            "degree": "Bachelor of Technology in Computer Science",
            "institution": "Apex Institute of Technology",
            "year": "2021 - 2025",
            "gpa": "3.8 / 4.0"
        }
    ],
    "certifications": [
        {
            "name": "Meta Front-End Developer Professional Certificate",
            "issuer": "Coursera",
            "issue_date": "2024"
        },
        {
            "name": "Python for Data Science and Machine Learning",
            "issuer": "Udemy",
            "issue_date": "2023"
        }
    ],
    "experience": [
        {
            "role": "Frontend Web Development Intern",
            "company": "Apex Web Solutions",
            "duration": "June 2024 - August 2024",
            "description": "Collaborated with agile engineering team to implement responsive landing pages using HTML, CSS, and JavaScript. Improved web accessibility standards across 12 product pages."
        }
    ],
    "achievements": [
        "First Place Winner - Apex Hackathon 2024 (Team of 4)",
        "Completed 250+ algorithmic problem solutions on LeetCode (Rating: 1680)",
        "Dean's Honor List for Academic Excellence (Consecutive semesters: 2022-2024)"
    ]
}

SAMPLE_FRONTEND_JD = """
Job Title: Junior Frontend Developer
Company: Horizon Tech Labs
Location: Remote / Hybrid

About the Role:
We are looking for an enthusiastic Frontend Developer to join our core web engineering team. You will build user-friendly single-page applications, implement responsive designs, and integrate REST APIs.

Key Responsibilities:
- Build clean, accessible, modern web interfaces using HTML, CSS, and JavaScript.
- Develop interactive user interface components with React.
- Collaborate with backend engineers to integrate REST API services.
- Utilize Git and GitHub for collaborative branching and pull request code reviews.
- Ensure cross-browser compatibility and responsive performance across mobile and desktop.

Requirements:
- Strong proficiency in HTML5, CSS3, and modern JavaScript (ES6+).
- Practical experience building web apps with React.
- Understanding of REST API integration and state handling.
- Familiarity with version control using Git.
- Demonstrated projects or code repositories on GitHub.
"""

SAMPLE_AIML_JD = """
Job Title: Associate AI/ML Engineer
Company: DataSphere Intelligence
Location: Remote

About the Role:
DataSphere Intelligence is seeking an Associate AI/ML Engineer to train predictive machine learning models, process complex datasets, and build evaluation pipelines.

Key Responsibilities:
- Design and train supervised machine learning models using Python and Scikit-Learn.
- Perform exploratory data analysis and feature engineering utilizing Pandas and NumPy.
- Build clean, reproducible data preprocessing workflows and model scoring pipelines.
- Collaborate with software engineers to package predictive models via REST API endpoints.

Requirements:
- Hands-on proficiency in Python programming for data science.
- Strong understanding of core Machine Learning algorithms (regression, classification, clustering).
- Experience with Scikit-Learn, Pandas, NumPy, and SQL.
- Strong problem-solving skills and verifiable projects demonstrating model development.
"""

def generate_sample_docx(data, output_path):
    doc = docx.Document()
    contact = data["contact"]

    title_p = doc.add_paragraph()
    r = title_p.add_run(contact["name"])
    r.bold = True
    r.font.size = docx.shared.Pt(18)
    title_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER

    contact_p = doc.add_paragraph()
    contact_p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
    contact_p.add_run(f"{contact['email']} | {contact['phone']} | {contact['github']} | {contact['linkedin']}")

    # Summary
    doc.add_heading("Professional Summary", level=1)
    doc.add_paragraph(data["summary"])

    # Skills
    doc.add_heading("Technical Skills", level=1)
    for cat, slist in data["skills"].items():
        p = doc.add_paragraph()
        p.add_run(f"{cat}: ").bold = True
        p.add_run(", ".join(slist))

    # Projects
    doc.add_heading("Projects", level=1)
    for proj in data["projects"]:
        p = doc.add_paragraph()
        p.add_run(f"{proj['title']} | {proj['technologies']}").bold = True
        doc.add_paragraph(proj["description"], style='List Bullet')
        if proj.get("github_url"):
            doc.add_paragraph(f"Repository: {proj['github_url']}", style='List Bullet')

    # Experience
    doc.add_heading("Experience", level=1)
    for exp in data["experience"]:
        p = doc.add_paragraph()
        p.add_run(f"{exp['role']} – {exp['company']} ({exp['duration']})").bold = True
        doc.add_paragraph(exp["description"], style='List Bullet')

    # Education
    doc.add_heading("Education", level=1)
    for edu in data["education"]:
        p = doc.add_paragraph()
        p.add_run(f"{edu['degree']} – {edu['institution']} ({edu['year']}) - GPA: {edu['gpa']}").bold = True

    # Certifications
    doc.add_heading("Certifications", level=1)
    for cert in data["certifications"]:
        doc.add_paragraph(f"{cert['name']} – {cert['issuer']} ({cert['issue_date']})", style='List Bullet')

    # Achievements
    doc.add_heading("Achievements", level=1)
    for ach in data["achievements"]:
        doc.add_paragraph(ach, style='List Bullet')

    doc.save(output_path)
    return output_path

def generate_sample_txt(data, output_path):
    contact = data["contact"]
    lines = [
        contact["name"],
        f"{contact['email']} | {contact['phone']} | {contact['github']} | {contact['linkedin']}",
        "",
        "SUMMARY",
        data["summary"],
        "",
        "EDUCATION",
        f"{data['education'][0]['degree']}",
        f"{data['education'][0]['institution']} | {data['education'][0]['year']} | GPA: {data['education'][0]['gpa']}",
        "",
        "SKILLS",
        ", ".join([s for cat in data["skills"].values() for s in cat]),
        "",
        "PROJECTS"
    ]
    for p in data["projects"]:
        lines.append(f"{p['title']} | {p['technologies']}")
        lines.append(p["description"])
        if p.get("github_url"):
            lines.append(f"GitHub: {p['github_url']}")
        lines.append("")

    lines.append("EXPERIENCE")
    for exp in data["experience"]:
        lines.append(f"{exp['role']} - {exp['company']} | {exp['duration']}")
        lines.append(exp["description"])
        lines.append("")

    lines.append("CERTIFICATIONS")
    for cert in data["certifications"]:
        lines.append(f"{cert['name']} - {cert['issuer']} ({cert['issue_date']})")
    lines.append("")

    lines.append("ACHIEVEMENTS")
    for ach in data["achievements"]:
        lines.append(f"• {ach}")

    content = "\n".join(lines)
    with open(output_path, "w", encoding="utf-8") as f:
        f.write(content)
    return output_path

if __name__ == "__main__":
    base_dir = os.path.dirname(os.path.abspath(__file__))
    pdf_path = os.path.join(base_dir, "sample_resume_alex_kumar.pdf")
    docx_path = os.path.join(base_dir, "sample_resume_alex_kumar.docx")
    txt_path = os.path.join(base_dir, "sample_resume_alex_kumar.txt")

    generate_pdf_resume(SAMPLE_RESUME_DATA, pdf_path)
    generate_sample_docx(SAMPLE_RESUME_DATA, docx_path)
    generate_sample_txt(SAMPLE_RESUME_DATA, txt_path)

    # Save sample JDs
    with open(os.path.join(base_dir, "sample_frontend_jd.txt"), "w", encoding="utf-8") as f:
        f.write(SAMPLE_FRONTEND_JD.strip())
    with open(os.path.join(base_dir, "sample_aiml_jd.txt"), "w", encoding="utf-8") as f:
        f.write(SAMPLE_AIML_JD.strip())

    print("Sample resumes (PDF, DOCX, TXT) and Job Descriptions generated successfully!")
